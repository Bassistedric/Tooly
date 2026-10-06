from __future__ import annotations

import argparse
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

from app.db.session import SessionLocal
from app.domains.inspections import service
from app.domains.inspections.schemas import (
    TemplateCheckpointCreate,
    TemplateCreate,
    TemplateSectionCreate,
)


def stable_code(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii").upper()
    return re.sub(r"[^A-Z0-9]+", "_", ascii_value).strip("_")


def parse_sheet(sheet, source_name: str) -> TemplateCreate:
    title = str(sheet["B2"].value or sheet.title).strip()

    header_row = None
    for row in range(1, sheet.max_row + 1):
        if (
            sheet.cell(row, 2).value == "CATÉGORIE"
            and sheet.cell(row, 3).value == "POINT DE VÉRIFICATION"
        ):
            header_row = row
            break

    if header_row is None:
        raise ValueError(f"No checkpoint table found in sheet: {sheet.title}")

    reminders: list[str] = []
    for row in range(13, header_row):
        value = sheet.cell(row, 2).value
        if isinstance(value, str) and value.strip().startswith("•"):
            reminders.append(value.strip())

    sections: list[TemplateSectionCreate] = []
    current_title: str | None = None
    current_points: list[TemplateCheckpointCreate] = []

    def flush_section() -> None:
        nonlocal current_title, current_points
        if current_title and current_points:
            sections.append(
                TemplateSectionCreate(
                    title=current_title,
                    checkpoints=current_points,
                )
            )
        current_title = None
        current_points = []

    for row in range(header_row + 1, sheet.max_row + 1):
        category = sheet.cell(row, 2).value
        checkpoint = sheet.cell(row, 3).value
        if not checkpoint:
            continue

        if category:
            flush_section()
            current_title = str(category).strip()
        elif current_title is None:
            current_title = "Général"

        current_points.append(
            TemplateCheckpointCreate(
                text=str(checkpoint).strip(),
                allows_na=True,
                is_required=True,
            )
        )

    flush_section()

    if not sections:
        raise ValueError(f"No checkpoints found in sheet: {sheet.title}")

    return TemplateCreate(
        code=stable_code(sheet.title),
        name=sheet.title,
        title=title,
        reminders="\n".join(reminders) or None,
        source_reference=f"{source_name} :: {sheet.title}",
        sections=sections,
    )


def import_workbook(path: Path) -> None:
    workbook = load_workbook(path, data_only=False, read_only=True)
    db = SessionLocal()
    created = 0
    skipped = 0

    try:
        for sheet in workbook.worksheets:
            if sheet.title == "Accueil":
                continue

            data = parse_sheet(sheet, path.name)
            try:
                service.create_template(db, data)
                created += 1
                print(f"CREATED {data.code}")
            except service.InspectionTemplateConflictError:
                skipped += 1
                print(f"SKIPPED {data.code} (already exists)")

        print(f"Done: {created} created, {skipped} skipped")
    finally:
        db.close()
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    import_workbook(args.workbook)


if __name__ == "__main__":
    main()
